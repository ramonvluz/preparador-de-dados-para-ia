import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from preparador_dados_ia.converters.base import ConversionContext
from preparador_dados_ia.converters.xlsx import XlsxWorkbookConverter
from preparador_dados_ia.converters.xlsx.extractor import InvalidXlsxError
from preparador_dados_ia.core.cancellation import CancellationToken
from preparador_dados_ia.core.cleaning import UnicodeSanitizer

PROJECT_ROOT = Path(__file__).parents[2]
SCHEMA = json.loads(
    (PROJECT_ROOT / "schemas" / "spreadsheet_workbook.schema.json").read_text(encoding="utf-8")
)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Vendas"
    sales.append(["Produto", "Quantidade", "Preço", "Total", "Data"])
    sales.append(["Caderno", 2, 12.5, "=B2*C2", datetime(2026, 8, 1)])
    sales.append(["Caneta", 3, 4.0, "=B3*C3", datetime(2026, 8, 2)])
    table = Table(displayName="TabelaVendas", ref="A1:E3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sales.add_table(table)
    sales.row_dimensions[3].hidden = True
    sales.column_dimensions["E"].hidden = True

    notes = workbook.create_sheet("Resumo")
    notes.append(["Indicador", "Valor"])
    notes.append(["Itens", 2])
    notes.merge_cells("A4:B4")
    notes["A4"] = "Observação"
    notes.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def _context(source: Path, **options: object) -> ConversionContext:
    return ConversionContext(
        source=source,
        source_size_bytes=source.stat().st_size,
        converted_at=datetime(2026, 8, 19, 12, 0, tzinfo=UTC),
        cancellation_token=CancellationToken(),
        unicode_sanitizer=UnicodeSanitizer(),
        options=options,
    )


def test_xlsx_converter_preserves_sheets_tables_types_and_formulas(tmp_path: Path) -> None:
    source = tmp_path / "vendas.xlsx"
    _workbook(source)

    records = list(XlsxWorkbookConverter().convert(_context(source)))

    assert len(records) == 2
    for record in records:
        Draft202012Validator(SCHEMA).validate(record)
    sales, summary = records
    assert sales["record_type"] == "spreadsheet_workbook"
    assert sales["data"]["sheet"]["name"] == "Vendas"
    assert sales["data"]["dataset"]["source_kind"] == "excel_table"
    assert sales["data"]["dataset"]["name"] == "TabelaVendas"
    assert sales["data"]["dataset"]["rows"][0][:3] == ["Caderno", 2, 12.5]
    assert sales["data"]["formulas"][0]["formula"] == "=B2*C2"
    assert sales["data"]["formulas"][0]["cached_value"] is None
    assert sales["processing"]["external_links_followed"] is False
    assert summary["data"]["sheet"]["visibility"] == "hidden"
    assert summary["processing"]["merged_ranges"] == 1


def test_xlsx_converter_segments_each_dataset_without_losing_formula_ranges(
    tmp_path: Path,
) -> None:
    source = tmp_path / "vendas.xlsx"
    _workbook(source)

    records = list(XlsxWorkbookConverter().convert(_context(source, xlsx_rows_per_segment=1)))
    sales = [record for record in records if record["data"]["sheet"]["name"] == "Vendas"]

    assert len(sales) == 2
    assert [record["data"]["dataset"]["row_start"] for record in sales] == [1, 2]
    assert [record["data"]["formulas"][0]["cell_reference"] for record in sales] == [
        "D2",
        "D3",
    ]


def test_xlsx_converter_rejects_invalid_zip_package(tmp_path: Path) -> None:
    source = tmp_path / "invalido.xlsx"
    source.write_text("não é uma planilha", encoding="utf-8")

    with pytest.raises(InvalidXlsxError, match="pacote XLSX válido"):
        list(XlsxWorkbookConverter().convert(_context(source)))
