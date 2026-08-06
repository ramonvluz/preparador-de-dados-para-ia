from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).parents[2]
FIXTURE = PROJECT_ROOT / "tests" / "fixtures" / "artificial_document.docx"


def test_word_fixture_matches_standard_business_brief_geometry() -> None:
    document = Document(FIXTURE)
    section = document.sections[0]

    assert section.page_width.inches == pytest.approx(8.5)
    assert section.page_height.inches == pytest.approx(11)
    assert section.top_margin.inches == pytest.approx(1.0)
    assert section.right_margin.inches == pytest.approx(1.0)
    assert section.bottom_margin.inches == pytest.approx(1.0)
    assert section.left_margin.inches == pytest.approx(1.0)
    assert section.header_distance.inches == pytest.approx(0.492, abs=0.001)
    assert section.footer_distance.inches == pytest.approx(0.492, abs=0.001)

    normal = document.styles["Normal"]
    assert normal.font.name == "Calibri"
    assert normal.font.size.pt == pytest.approx(11)
    assert normal.paragraph_format.space_after.pt == pytest.approx(6)
    assert normal.paragraph_format.line_spacing == pytest.approx(1.10)
    assert section.header.paragraphs[0].text == "LIMEBH | FIXTURE DOCX 3D"
    assert "nenhum dado real" in section.footer.paragraphs[0].text


def test_word_fixture_uses_real_lists_and_fixed_table_geometry() -> None:
    document = Document(FIXTURE)
    list_paragraphs = [
        paragraph for paragraph in document.paragraphs if paragraph.style.name.startswith("List")
    ]
    assert len(list_paragraphs) == 5
    assert all(paragraph.style.element.pPr.numPr is not None for paragraph in list_paragraphs)

    table = document.tables[0]
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    indent = properties.find(qn("w:tblInd"))
    layout = properties.find(qn("w:tblLayout"))
    assert width.get(qn("w:w")) == "9360"
    assert width.get(qn("w:type")) == "dxa"
    assert indent.get(qn("w:w")) == "120"
    assert layout.get(qn("w:type")) == "fixed"
    assert [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst] == [
        2160,
        3600,
        3600,
    ]
    assert all(
        cell._tc.get_or_add_tcPr().tcW.get(qn("w:type")) == "dxa"
        for row in table.rows
        for cell in row.cells
    )
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(row.height is None for row in table.rows)
