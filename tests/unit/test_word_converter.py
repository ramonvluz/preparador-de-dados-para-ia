import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from jsonschema import Draft202012Validator

from limebh_preparador.application.progress import ConversionProgress
from limebh_preparador.contracts.word import WordParagraph
from limebh_preparador.converters.base import ConversionContext
from limebh_preparador.converters.word import WordDocumentConverter
from limebh_preparador.converters.word.extractor import suspicious_text_sequence_count
from limebh_preparador.core.cancellation import CancellationToken

PROJECT_ROOT = Path(__file__).parents[2]
FIXTURE = PROJECT_ROOT / "tests" / "fixtures" / "artificial_document.docx"
SCHEMA = PROJECT_ROOT / "schemas" / "word_document.schema.json"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_suspicious_text_detector_avoids_common_legitimate_symbols() -> None:
    legitimate = WordParagraph("O 1º lugar usa ½ porção e precisão de 50 µm.")
    corrupted = WordParagraph("InformaÃ§Ã£o com Definiçºµo corrompida.")

    assert suspicious_text_sequence_count([legitimate]) == 0
    assert suspicious_text_sequence_count([corrupted]) == 3


def test_word_converter_extracts_ordered_valid_contract_without_changing_source() -> None:
    source_hash = _sha256(FIXTURE)
    progress: list[ConversionProgress] = []
    context = ConversionContext(
        source=FIXTURE,
        source_size_bytes=FIXTURE.stat().st_size,
        converted_at=datetime(2026, 8, 6, 15, 30, tzinfo=UTC),
        cancellation_token=CancellationToken(),
        progress_callback=progress.append,
    )

    records = list(WordDocumentConverter().convert(context))

    assert _sha256(FIXTURE) == source_hash
    assert len(records) == 1
    record = records[0]
    Draft202012Validator(json.loads(SCHEMA.read_text(encoding="utf-8"))).validate(record)
    assert record["source"]["file_name"] == "artificial_document.docx"
    assert record["data"]["title"] == "Guia Artificial de Validação 3D"
    assert record["data"]["author"] == "LIMEBH"
    assert len(record["data"]["blocks"]) == 18
    assert [block["type"] for block in record["data"]["blocks"]][3:7] == [
        "heading",
        "paragraph",
        "heading",
        "list_item",
    ]
    list_items = [block for block in record["data"]["blocks"] if block["type"] == "list_item"]
    assert [(item["level"], item["ordered"]) for item in list_items] == [
        (0, False),
        (0, False),
        (1, False),
        (0, True),
        (0, True),
    ]
    table = next(block for block in record["data"]["blocks"] if block["type"] == "table")
    assert table["has_header"] is True
    assert table["rows"][3][2] == "Risco | mitigação"
    assert len(record["data"]["sections"]) == 5
    assert record["data"]["sections"][1] == {
        "title": "1.1 Critérios",
        "level": 2,
        "start_block": 5,
        "end_block": 10,
    }
    assert record["processing"]["images_omitted"] == 1
    assert record["processing"]["headers_footers_omitted"] is True
    assert record["processing"]["features_omitted"] == []
    assert record["processing"]["suspicious_text_sequences"] == 0
    assert record["processing"]["inferred_headings"] == 0
    assert record["processing"]["inferred_table_headers"] == 0
    assert progress[0].current == 1
    assert progress[-1].current == progress[-1].total


def test_word_converter_warns_about_suspicious_text_and_infers_safe_structure(
    tmp_path: Path,
) -> None:
    source = tmp_path / "sem_estilos.docx"
    document = Document()
    title = document.add_paragraph()
    title_run = title.add_run("Documento sem estilos")
    title_run.bold = True
    title_run.font.size = Pt(18)
    heading = document.add_paragraph()
    heading_run = heading.add_run("1. Definiçºµo")
    heading_run.bold = True
    heading_run.font.size = Pt(12)
    document.add_paragraph("Texto preservado exatamente como aparece na origem.")
    table = document.add_table(rows=2, cols=3)
    table_look = table._tbl.tblPr.find(qn("w:tblLook"))
    if table_look is not None:
        table_look.set(qn("w:firstRow"), "0")
    for cell, value in zip(
        table.rows[0].cells,
        ("Serviço", "Objetivo", "Resultado"),
        strict=True,
    ):
        cell.text = value
    for cell, value in zip(
        table.rows[1].cells,
        (
            "Diagnóstico Estratégico",
            "Compreender detalhadamente a operação da empresa",
            "Plano executivo priorizado para orientar decisões",
        ),
        strict=True,
    ):
        cell.text = value
    document.save(source)
    context = ConversionContext(
        source=source,
        source_size_bytes=source.stat().st_size,
        converted_at=datetime(2026, 8, 6, 18, 30, tzinfo=UTC),
        cancellation_token=CancellationToken(),
    )

    record = next(WordDocumentConverter().convert(context))

    Draft202012Validator(json.loads(SCHEMA.read_text(encoding="utf-8"))).validate(record)
    assert record["data"]["title"] == "Documento sem estilos"
    headings = [block for block in record["data"]["blocks"] if block["type"] == "heading"]
    assert [(block["text"], block["level"], block["inferred"]) for block in headings] == [
        ("Documento sem estilos", 1, True),
        ("1. Definiçºµo", 1, True),
    ]
    table_block = next(block for block in record["data"]["blocks"] if block["type"] == "table")
    assert table_block["has_header"] is True
    assert table_block["header_inferred"] is True
    assert record["processing"]["suspicious_text_sequences"] == 1
    assert record["processing"]["inferred_headings"] == 2
    assert record["processing"]["inferred_table_headers"] == 1
    assert record["processing"]["warnings"] == ["word_suspected_text_encoding_corruption:1"]
