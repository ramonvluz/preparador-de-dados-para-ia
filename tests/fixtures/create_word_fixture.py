"""Gera o DOCX artificial usado nos testes da fase 3D."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

OUTPUT = Path(__file__).with_name("artificial_document.docx")
FIXED_TIME = datetime(2026, 8, 6, 15, 0, 0)
INK = RGBColor(0x18, 0x33, 0x2E)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)


def _set_font(run: object, *, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _configure_styles(document: object) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)


def _configure_section(section: object) -> None:
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    _set_font(
        header.add_run("PREPARADOR | FIXTURE DOCX 3D"),
        size=8,
        color=RGBColor(0x61, 0x73, 0x6F),
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    _set_font(
        footer.add_run("Documento artificial - nenhum dado real"),
        size=8,
        color=RGBColor(0x61, 0x73, 0x6F),
    )


def _add_bottom_rule(paragraph: object) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "146B5C")
    borders.append(bottom)
    properties.append(borders)


def _set_cell_margins(table: object) -> None:
    properties = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)


def _configure_table(table: object, widths: tuple[int, ...]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    _set_cell_margins(table)

    grid_columns = table._tbl.tblGrid.gridCol_lst
    for column, column_width in zip(grid_columns, widths, strict=True):
        column.set(qn("w:w"), str(column_width))
    for row in table.rows:
        for cell, cell_width in zip(row.cells, widths, strict=True):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(cell_width / 1440)
            cell._tc.get_or_add_tcPr().tcW.set(qn("w:w"), str(cell_width))
            cell._tc.get_or_add_tcPr().tcW.set(qn("w:type"), "dxa")

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_properties.append(OxmlElement("w:tblHeader"))
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F4F7")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True


def _status_image() -> BytesIO:
    image = Image.new("RGB", (600, 170), "#eef5f2")
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((8, 8, 592, 162), outline="#146b5c", width=5)
    drawing.text((45, 55), "IMAGEM ARTIFICIAL - CONTEUDO NAO EXTRAIDO", fill="#18332e")
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def _repack_deterministically(source: bytes, output: Path) -> None:
    with ZipFile(BytesIO(source)) as archive, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            info = ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            target.writestr(info, archive.read(name))


def create_fixture(output: Path = OUTPUT) -> Path:
    document = Document()
    document.core_properties.title = "Guia Artificial de Validação 3D"
    document.core_properties.author = "Projeto Artificial"
    document.core_properties.subject = "Fixture DOCX artificial"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.core_properties.last_modified_by = "Projeto Artificial"
    _configure_styles(document)
    _configure_section(document.sections[0])

    document.add_paragraph("Guia Artificial de Validação 3D", style="Title")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(3)
    _set_font(subtitle.add_run("Contrato word_document e saída Markdown"), size=13, color=DARK_BLUE)
    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(14)
    _set_font(metadata.add_run("Status: "), size=10, bold=True)
    _set_font(metadata.add_run("Amostra aprovada para testes locais"), size=10)
    _add_bottom_rule(metadata)

    document.add_heading("1. Objetivo", level=1)
    document.add_paragraph(
        "Validar a preservação de títulos, parágrafos, listas, seções e tabelas simples "
        "sem alterar o documento de origem."
    )
    document.add_heading("1.1 Critérios", level=2)
    document.add_paragraph(
        "Preservar a ordem dos blocos e a acentuação em português.", style="List Bullet"
    )
    document.add_paragraph("Manter listas aninhadas como estrutura Markdown.", style="List Bullet")
    document.add_paragraph("Subitem artificial para validar nível.", style="List Bullet 2")
    document.add_paragraph("Gerar partes independentes dentro dos limites.", style="List Number")
    document.add_paragraph("Registrar recursos omitidos no relatório.", style="List Number")

    document.add_heading("2. Matriz de validação", level=1)
    document.add_paragraph("A tabela abaixo contém somente informações artificiais.")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(
        table.rows[0].cells, ("Elemento", "Expectativa", "Resultado"), strict=True
    ):
        cell.text = value
    for values in (
        ("Título", "Hierarquia preservada", "Markdown com #"),
        ("Lista", "Ordem e nível", "Marcadores válidos"),
        ("Tabela", "Células e cabeçalho", "Risco | mitigação"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    _configure_table(table, (2160, 3600, 3600))

    document.add_heading("3. Elemento visual", level=1)
    document.add_paragraph(
        "A imagem a seguir deve ser contabilizada como omitida, sem copiar seu conteúdo binário."
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(_status_image(), width=Inches(4.2))

    document.add_page_break()
    document.add_heading("4. Conclusão", level=1)
    document.add_paragraph(
        "O documento artificial termina com uma segunda página para validar transições, "
        "metadados e leitura completa do pacote DOCX."
    )

    buffer = BytesIO()
    document.save(buffer)
    _repack_deterministically(buffer.getvalue(), output)
    return output


if __name__ == "__main__":
    print(create_fixture())
