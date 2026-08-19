import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from preparador_dados_ia.application.progress import ConversionProgress
from preparador_dados_ia.application.service import convert_source
from preparador_dados_ia.core.cancellation import CancellationToken

PROJECT_ROOT = Path(__file__).parents[2]
FIXTURE = PROJECT_ROOT / "tests" / "fixtures" / "artificial_document.docx"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_word_conversion_generates_structured_markdown_and_audit_report(
    tmp_path: Path,
) -> None:
    source_hash = _sha256(FIXTURE)
    output = tmp_path / "word_saida"

    report = convert_source(FIXTURE, output)

    assert _sha256(FIXTURE) == source_hash
    assert report["result"] == "success_with_warnings"
    assert report["converted_records"] == 1
    assert report["failed_records"] == 0
    assert report["word_extraction"] == {
        "block_count": 18,
        "paragraphs": 7,
        "headings": 5,
        "list_items": 5,
        "tables": 1,
        "sections": 5,
        "inferred_headings": 0,
        "inferred_table_headers": 0,
    }
    assert report["text_quality"] == {
        "suspected_encoding_corruption": False,
        "suspicious_sequences": 0,
        "automatic_text_repair_applied": False,
    }
    assert report["omitted_content"] == {
        "images": 1,
        "headers_footers": True,
        "advanced_features": [],
        "external_references_downloaded": False,
    }
    ready_files = list((output / "PRONTO_PARA_IA").glob("*.md"))
    assert ready_files
    ready_text = "\n".join(path.read_text(encoding="utf-8") for path in ready_files)
    assert "Guia Artificial de Validação 3D" in ready_text
    assert "# 1. Objetivo" in ready_text
    assert "    - Subitem artificial para validar nível." in ready_text
    assert "Risco \\| mitigação" in ready_text
    assert "IMAGEM ARTIFICIAL" not in ready_text
    assert "1 imagem(ns) não extraída(s)" in ready_text
    assert "PREPARADOR_" not in ready_text
    assert "schema_version" not in ready_text
    assert str(FIXTURE.resolve()) not in ready_text
    readme = (output / "LEIA-ME.txt").read_text(encoding="utf-8")
    assert "Tipo identificado: DOCX" in readme
    assert "Títulos inferidos por formatação: 0" in readme
    assert (output / "relatorio_conversao.json").is_file()


def test_word_conversion_surfaces_quality_warning_and_inferred_structure(
    tmp_path: Path,
) -> None:
    source = tmp_path / "origem_sem_estilos.docx"
    document = Document()
    title = document.add_paragraph()
    title_run = title.add_run("Documento sem estilos")
    title_run.bold = True
    title_run.font.size = Pt(18)
    heading = document.add_paragraph()
    heading_run = heading.add_run("1. Definiçºµo")
    heading_run.bold = True
    heading_run.font.size = Pt(12)
    document.add_paragraph("Conteúdo de validação preservado.")
    table = document.add_table(rows=2, cols=3)
    table_look = table._tbl.tblPr.find(qn("w:tblLook"))
    if table_look is not None:
        table_look.set(qn("w:firstRow"), "0")
    for cell, value in zip(
        table.rows[0].cells,
        ("Serviço", "Objetivo", "Resultado"),
        strict=True,
    ):
        cell.text = value
    for cell, value in zip(
        table.rows[1].cells,
        (
            "Diagnóstico Estratégico",
            "Compreender detalhadamente a operação da empresa",
            "Plano executivo priorizado para orientar decisões",
        ),
        strict=True,
    ):
        cell.text = value
    document.save(source)
    output = tmp_path / "saida_sem_estilos"

    report = convert_source(source, output)

    assert report["result"] == "success_with_warnings"
    assert report["warnings_count"] == 1
    assert report["word_extraction"] == {
        "block_count": 4,
        "paragraphs": 1,
        "headings": 2,
        "list_items": 0,
        "tables": 1,
        "sections": 2,
        "inferred_headings": 2,
        "inferred_table_headers": 1,
    }
    assert report["text_quality"] == {
        "suspected_encoding_corruption": True,
        "suspicious_sequences": 1,
        "automatic_text_repair_applied": False,
    }
    markdown = (output / "PRONTO_PARA_IA" / "documento_parte_0001.md").read_text(encoding="utf-8")
    assert "# Documento sem estilos" in markdown
    assert "# 1. Definiçºµo" in markdown
    assert "| Serviço | Objetivo | Resultado |" in markdown
    assert "| Coluna 1 |" not in markdown
    assert "1 sequência(s) de texto suspeita(s)" in markdown
    readme = (output / "LEIA-ME.txt").read_text(encoding="utf-8")
    assert "Tipo identificado: DOCX" in readme
    assert "ATENÇÃO: foram detectadas 1 sequência(s)" in readme


def test_invalid_docx_fails_without_exposing_content(tmp_path: Path) -> None:
    source = tmp_path / "invalido.docx"
    source.write_text("segredo-artificial-fora-de-um-pacote", encoding="utf-8")
    output = tmp_path / "invalido_saida"

    report = convert_source(source, output)

    assert report["result"] == "failure"
    assert report["errors"][0]["error_type"] == "InvalidDocxError"
    report_text = (output / "relatorio_conversao.json").read_text(encoding="utf-8")
    assert "segredo-artificial-fora-de-um-pacote" not in report_text


def test_word_cancellation_keeps_only_complete_artifacts(tmp_path: Path) -> None:
    token = CancellationToken()

    def cancel_after_first_block(progress: ConversionProgress) -> None:
        if progress.stage == "converting" and progress.current == 1:
            token.cancel()

    output = tmp_path / "word_cancelado"
    report = convert_source(
        FIXTURE,
        output,
        cancellation_token=token,
        progress_callback=cancel_after_first_block,
    )

    assert report["result"] == "cancelled"
    assert report["converted_records"] == 0
    assert list(output.rglob("*.tmp")) == []
    assert (output / "LEIA-ME.txt").is_file()
    assert (output / "relatorio_conversao.json").is_file()
