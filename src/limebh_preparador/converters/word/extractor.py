from __future__ import annotations

import re
from collections.abc import Callable, Iterator
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from limebh_preparador.contracts.word import (
    WordBlock,
    WordHeading,
    WordListItem,
    WordParagraph,
    WordSection,
    WordTable,
)
from limebh_preparador.core.cancellation import CancellationToken
from limebh_preparador.core.cleaning import UnicodeSanitizer

MAX_DOCX_ENTRY_BYTES = 64 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 1_000
HEADING_STYLE = re.compile(r"^(?:Heading|Título)\s*([1-9])$", re.IGNORECASE)
HEADING_STYLE_ID = re.compile(r"^Heading([1-9])$", re.IGNORECASE)
LIST_LEVEL_SUFFIX = re.compile(r"([2-9])$")


class InvalidDocxError(ValueError):
    """Indica que a fonte não é um pacote DOCX seguro e legível."""


def validate_docx_archive(path: Path) -> set[str]:
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise InvalidDocxError("O arquivo não contém a estrutura obrigatória de um DOCX")
            total = 0
            for info in archive.infolist():
                total += info.file_size
                if info.file_size > MAX_DOCX_ENTRY_BYTES:
                    raise InvalidDocxError("O DOCX contém uma parte interna excessivamente grande")
                if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise InvalidDocxError(
                        "O conteúdo descompactado do DOCX excede o limite seguro"
                    )
                if (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise InvalidDocxError("O DOCX apresenta uma taxa de compactação suspeita")
            return names
    except BadZipFile as error:
        raise InvalidDocxError("O arquivo não é um pacote DOCX válido") from error


def extract_blocks(
    document: DocumentObject,
    sanitizer: UnicodeSanitizer,
    cancellation_token: CancellationToken,
    progress_callback: Callable[[int, int], None] | None = None,
) -> tuple[list[WordBlock], list[WordSection]]:
    elements = list(document.iter_inner_content())
    total = len(elements)
    blocks: list[WordBlock] = []
    for current, element in enumerate(elements, start=1):
        cancellation_token.raise_if_cancelled()
        block = _extract_block(element, sanitizer)
        if block is not None:
            blocks.append(block)
        if progress_callback is not None:
            progress_callback(current, total)
    return blocks, build_sections(blocks)


def _extract_block(
    element: Paragraph | Table,
    sanitizer: UnicodeSanitizer,
) -> WordBlock | None:
    if isinstance(element, Table):
        rows = tuple(
            tuple(_cell_text(cell, sanitizer) for cell in row.cells) for row in element.rows
        )
        if not rows or not any(any(cell for cell in row) for row in rows):
            return None
        return WordTable(rows=rows, has_header=_table_has_header(element))

    text = sanitizer.clean(element.text).strip()
    if not text:
        return None
    heading_level = _heading_level(element)
    if heading_level is not None:
        return WordHeading(text=text, level=heading_level)
    list_info = _list_info(element)
    if list_info is not None:
        level, ordered = list_info
        return WordListItem(text=text, level=level, ordered=ordered)
    style_name = element.style.name if element.style is not None else None
    return WordParagraph(text=text, style=style_name)


def _cell_text(cell: object, sanitizer: UnicodeSanitizer) -> str:
    paragraphs = getattr(cell, "paragraphs", ())
    values = [sanitizer.clean(paragraph.text).strip() for paragraph in paragraphs]
    return "\n".join(value for value in values if value)


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    if style is not None:
        style_id_match = HEADING_STYLE_ID.match(style.style_id or "")
        name_match = HEADING_STYLE.match(style.name or "")
        match = style_id_match or name_match
        if match is not None:
            return int(match.group(1))
    outline_level = _outline_level(paragraph)
    return outline_level + 1 if outline_level is not None and outline_level < 9 else None


def _outline_level(paragraph: Paragraph) -> int | None:
    for properties in _paragraph_properties(paragraph):
        outline = properties.find(qn("w:outlineLvl"))
        if outline is None:
            continue
        value = outline.get(qn("w:val"))
        if value is not None and value.isdigit():
            return int(value)
    return None


def _list_info(paragraph: Paragraph) -> tuple[int, bool] | None:
    direct = _numbering_info(paragraph, paragraph._p.pPr)
    if direct is not None:
        return direct

    style = paragraph.style
    style_name = (style.name if style is not None else "") or ""
    style_id = (style.style_id if style is not None else "") or ""
    label = f"{style_name} {style_id}".lower()
    if "list" in label or "lista" in label:
        if "bullet" in label or "marcador" in label:
            ordered = False
        elif "number" in label or "número" in label or "numero" in label:
            ordered = True
        else:
            ordered = None
        if ordered is not None:
            suffix = LIST_LEVEL_SUFFIX.search(f"{style_name} {style_id}")
            return (int(suffix.group(1)) - 1 if suffix else 0), ordered

    style = paragraph.style
    visited: set[str] = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        inherited = _numbering_info(paragraph, style.element.pPr)
        if inherited is not None:
            return inherited
        style = style.base_style
    return None


def _numbering_info(paragraph: Paragraph, properties: object | None) -> tuple[int, bool] | None:
    if properties is None:
        return None
    num_properties = properties.find(qn("w:numPr"))  # type: ignore[attr-defined]
    if num_properties is None:
        return None
    level_element = num_properties.find(qn("w:ilvl"))
    number_element = num_properties.find(qn("w:numId"))
    level = _integer_attribute(level_element, "w:val", default=0) or 0
    number_id = _integer_attribute(number_element, "w:val", default=None)
    return level, _numbering_is_ordered(paragraph, number_id, level)


def _paragraph_properties(paragraph: Paragraph) -> Iterator[object]:
    if paragraph._p.pPr is not None:
        yield paragraph._p.pPr
    style = paragraph.style
    visited: set[str] = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        if style.element.pPr is not None:
            yield style.element.pPr
        style = style.base_style


def _numbering_is_ordered(paragraph: Paragraph, number_id: int | None, level: int) -> bool:
    if number_id is None:
        return True
    try:
        numbering = paragraph.part.numbering_part.element
        abstract_id: str | None = None
        for number in numbering.findall(qn("w:num")):
            if number.get(qn("w:numId")) != str(number_id):
                continue
            abstract = number.find(qn("w:abstractNumId"))
            abstract_id = abstract.get(qn("w:val")) if abstract is not None else None
            break
        if abstract_id is None:
            return True
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) != abstract_id:
                continue
            for level_element in abstract.findall(qn("w:lvl")):
                if level_element.get(qn("w:ilvl")) != str(level):
                    continue
                number_format = level_element.find(qn("w:numFmt"))
                value = number_format.get(qn("w:val")) if number_format is not None else None
                return value != "bullet"
    except (AttributeError, KeyError, ValueError):
        return True
    return True


def _integer_attribute(element: object, name: str, default: int | None) -> int | None:
    if element is None:
        return default
    value = element.get(qn(name))  # type: ignore[attr-defined]
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _table_has_header(table: Table) -> bool:
    if not table.rows:
        return False
    row_properties = table.rows[0]._tr.trPr
    if row_properties is not None and row_properties.find(qn("w:tblHeader")) is not None:
        return True
    table_properties = table._tbl.tblPr
    look = table_properties.find(qn("w:tblLook")) if table_properties is not None else None
    if look is not None and look.get(qn("w:firstRow")) in {"1", "true"}:
        return True
    return False


def build_sections(blocks: list[WordBlock]) -> list[WordSection]:
    sections: list[WordSection] = []
    for index, block in enumerate(blocks):
        if not isinstance(block, WordHeading):
            continue
        end_block = len(blocks) - 1
        for following_index in range(index + 1, len(blocks)):
            following = blocks[following_index]
            if isinstance(following, WordHeading) and following.level <= block.level:
                end_block = following_index - 1
                break
        sections.append(
            WordSection(
                title=block.text,
                level=block.level,
                start_block=index,
                end_block=max(index, end_block),
            )
        )
    return sections


def document_title(document: DocumentObject, sanitizer: UnicodeSanitizer) -> str | None:
    metadata_title = sanitizer.clean(document.core_properties.title or "").strip()
    if metadata_title:
        return metadata_title
    for paragraph in document.paragraphs:
        style_id = paragraph.style.style_id if paragraph.style is not None else ""
        if style_id == "Title" and paragraph.text.strip():
            return sanitizer.clean(paragraph.text).strip()
    for paragraph in document.paragraphs:
        if _heading_level(paragraph) == 1 and paragraph.text.strip():
            return sanitizer.clean(paragraph.text).strip()
    return None


def document_author(document: DocumentObject, sanitizer: UnicodeSanitizer) -> str | None:
    author = sanitizer.clean(document.core_properties.author or "").strip()
    return author or None


def embedded_image_count(document: DocumentObject) -> int:
    return sum(
        1 for part in document.part.package.parts if str(part.partname).startswith("/word/media/")
    )


def has_header_footer_content(document: DocumentObject) -> bool:
    seen: set[str] = set()
    for section in document.sections:
        for container in (section.header, section.footer):
            part_name = str(container.part.partname)
            if part_name in seen:
                continue
            seen.add(part_name)
            if any(paragraph.text.strip() for paragraph in container.paragraphs):
                return True
            if any(
                any(cell.text.strip() for row in table.rows for cell in row.cells)
                for table in container.tables
            ):
                return True
    return False


def omitted_features(archive_names: set[str], document: DocumentObject) -> list[str]:
    features: list[str] = []
    mapping = {
        "word/footnotes.xml": "footnotes",
        "word/endnotes.xml": "endnotes",
        "word/comments.xml": "comments",
    }
    for part_name, feature in mapping.items():
        if part_name in archive_names:
            features.append(feature)
    prefix_mapping = {
        "word/charts/": "charts",
        "word/embeddings/": "embedded_objects",
        "word/diagrams/": "diagrams",
    }
    for prefix, feature in prefix_mapping.items():
        if any(name.startswith(prefix) for name in archive_names):
            features.append(feature)
    if any(name.endswith("vbaProject.bin") for name in archive_names):
        features.append("macros")
    if any(cell.tables for table in document.tables for row in table.rows for cell in row.cells):
        features.append("nested_tables")
    xml = document.element.xml
    if "<w:ins" in xml or "<w:del" in xml:
        features.append("tracked_changes")
    if "<w:txbxContent" in xml:
        features.append("text_boxes")
    if "<m:oMath" in xml:
        features.append("equations")
    return features
