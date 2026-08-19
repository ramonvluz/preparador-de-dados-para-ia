from __future__ import annotations

import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter, range_boundaries

from limebh_preparador.contracts.tabular import (
    SpreadsheetFormula,
    SpreadsheetSheet,
    TabularColumn,
    TabularValue,
)
from limebh_preparador.core.cancellation import CancellationToken
from limebh_preparador.core.cleaning import UnicodeSanitizer

MAX_XLSX_ENTRY_BYTES = 64 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1_000
REQUIRED_XLSX_PARTS = frozenset({"[Content_Types].xml", "xl/workbook.xml"})


class InvalidXlsxError(ValueError):
    """Indica que a fonte não é um pacote XLSX seguro e legível."""


@dataclass(frozen=True, slots=True)
class ExtractedDataset:
    sheet: SpreadsheetSheet
    name: str
    source_kind: str
    reference: str
    has_header: bool
    columns: tuple[TabularColumn, ...]
    rows: tuple[tuple[TabularValue, ...], ...]
    formulas: tuple[tuple[int, SpreadsheetFormula], ...]
    warnings: tuple[str, ...]
    merged_ranges: int
    hidden_rows: int
    hidden_columns: int


@dataclass(frozen=True, slots=True)
class ExtractedWorkbook:
    sheet_count: int
    datasets: tuple[ExtractedDataset, ...]


def extract_workbook(
    path: Path,
    cancellation_token: CancellationToken,
    sanitizer: UnicodeSanitizer,
) -> ExtractedWorkbook:
    external_link_count = validate_xlsx_archive(path)
    formula_book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    value_book = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        datasets: list[ExtractedDataset] = []
        sheet_count = len(formula_book.worksheets)
        for sheet_index, formula_sheet in enumerate(formula_book.worksheets, start=1):
            cancellation_token.raise_if_cancelled()
            value_sheet = value_book[formula_sheet.title]
            sheet = SpreadsheetSheet(
                index=sheet_index,
                name=sanitizer.clean(formula_sheet.title),
                visibility=_sheet_visibility(formula_sheet.sheet_state),
                used_range=_used_range(formula_sheet),
            )
            common_warnings: list[str] = []
            if sheet.visibility != "visible":
                common_warnings.append(f"xlsx_hidden_sheet_included:{sheet.name}")
            if external_link_count:
                common_warnings.append(f"xlsx_external_links_not_followed:{external_link_count}")

            tables = list(formula_sheet.tables.values())
            if tables:
                for table in tables:
                    datasets.append(
                        _extract_range(
                            formula_sheet,
                            value_sheet,
                            sheet=sheet,
                            name=sanitizer.clean(str(table.displayName)),
                            source_kind="excel_table",
                            reference=str(table.ref),
                            has_header=True,
                            totals_rows=int(table.totalsRowCount or 0),
                            common_warnings=common_warnings,
                            sanitizer=sanitizer,
                            cancellation_token=cancellation_token,
                        )
                    )
            elif sheet.used_range is not None:
                bounds = range_boundaries(sheet.used_range)
                raw_first = [
                    formula_sheet.cell(row=bounds[1], column=column).value
                    for column in range(bounds[0], bounds[2] + 1)
                ]
                datasets.append(
                    _extract_range(
                        formula_sheet,
                        value_sheet,
                        sheet=sheet,
                        name=sheet.name,
                        source_kind="worksheet_range",
                        reference=sheet.used_range,
                        has_header=_looks_like_header(raw_first),
                        totals_rows=0,
                        common_warnings=common_warnings,
                        sanitizer=sanitizer,
                        cancellation_token=cancellation_token,
                    )
                )
            else:
                datasets.append(
                    _empty_dataset(
                        sheet,
                        warnings=(*common_warnings, "xlsx_empty_sheet"),
                    )
                )
        return ExtractedWorkbook(sheet_count=sheet_count, datasets=tuple(datasets))
    finally:
        formula_book.close()
        value_book.close()


def validate_xlsx_archive(path: Path) -> int:
    total_size = 0
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if not REQUIRED_XLSX_PARTS.issubset(names):
                raise InvalidXlsxError("O arquivo não contém a estrutura obrigatória de um XLSX")
            for info in archive.infolist():
                if info.flag_bits & 0x1:
                    raise InvalidXlsxError("O XLSX contém uma parte interna criptografada")
                total_size += info.file_size
                if info.file_size > MAX_XLSX_ENTRY_BYTES:
                    raise InvalidXlsxError("O XLSX contém uma parte interna excessivamente grande")
                if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise InvalidXlsxError(
                        "O conteúdo descompactado do XLSX excede o limite seguro"
                    )
                if (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise InvalidXlsxError("O XLSX apresenta uma taxa de compactação suspeita")
            return sum(1 for name in names if name.startswith("xl/externalLinks/externalLink"))
    except zipfile.BadZipFile as error:
        raise InvalidXlsxError("O arquivo não é um pacote XLSX válido") from error


def _extract_range(
    formula_sheet,
    value_sheet,
    *,
    sheet: SpreadsheetSheet,
    name: str,
    source_kind: str,
    reference: str,
    has_header: bool,
    totals_rows: int,
    common_warnings: list[str],
    sanitizer: UnicodeSanitizer,
    cancellation_token: CancellationToken,
) -> ExtractedDataset:
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    header_row = min_row if has_header else None
    data_start = min_row + 1 if has_header else min_row
    data_end = max(data_start - 1, max_row - totals_rows)
    width = max_col - min_col + 1
    raw_headers = [
        formula_sheet.cell(row=header_row, column=column).value if header_row else None
        for column in range(min_col, max_col + 1)
    ]
    source_headers = [
        sanitizer.clean(str(value)).strip() if value is not None else None for value in raw_headers
    ]
    names = _normalized_column_names(source_headers, width)
    rows: list[list[TabularValue]] = []
    formula_entries: list[tuple[int, SpreadsheetFormula]] = []
    for logical_row, worksheet_row in enumerate(range(data_start, data_end + 1), start=1):
        cancellation_token.raise_if_cancelled()
        output_row: list[TabularValue] = []
        for column in range(min_col, max_col + 1):
            formula_cell = formula_sheet.cell(row=worksheet_row, column=column)
            value_cell = value_sheet.cell(row=worksheet_row, column=column)
            cached_value = _normalize_cell_value(value_cell.value, sanitizer)
            output_row.append(cached_value)
            if not isinstance(formula_cell, MergedCell) and formula_cell.data_type == "f":
                formula = str(formula_cell.value or "")
                if not formula.startswith("="):
                    formula = f"={formula}"
                formula_entries.append(
                    (
                        logical_row,
                        SpreadsheetFormula(
                            cell_reference=formula_cell.coordinate,
                            formula=sanitizer.clean(formula),
                            cached_value=cached_value,
                        ),
                    )
                )
        rows.append(output_row)

    column_types = [_infer_column_type([row[index] for row in rows]) for index in range(width)]
    columns = tuple(
        TabularColumn(
            index=index + 1,
            name=names[index],
            source_name=source_headers[index],
            inferred_type=column_types[index],
            nullable=any(row[index] is None for row in rows),
        )
        for index in range(width)
    )
    formulas_without_cache = sum(
        1 for _, formula in formula_entries if formula.cached_value is None
    )
    warnings = list(common_warnings)
    if formulas_without_cache:
        warnings.append(f"xlsx_formulas_without_cached_value:{formulas_without_cache}")
    if totals_rows:
        warnings.append(f"xlsx_table_total_rows_omitted:{totals_rows}")
    return ExtractedDataset(
        sheet=sheet,
        name=name,
        source_kind=source_kind,
        reference=reference,
        has_header=has_header,
        columns=columns,
        rows=tuple(tuple(row) for row in rows),
        formulas=tuple(formula_entries),
        warnings=tuple(warnings),
        merged_ranges=len(formula_sheet.merged_cells.ranges),
        hidden_rows=sum(
            1 for dimension in formula_sheet.row_dimensions.values() if dimension.hidden
        ),
        hidden_columns=sum(
            1 for dimension in formula_sheet.column_dimensions.values() if dimension.hidden
        ),
    )


def _empty_dataset(
    sheet: SpreadsheetSheet,
    *,
    warnings: tuple[str, ...],
) -> ExtractedDataset:
    return ExtractedDataset(
        sheet=sheet,
        name=sheet.name,
        source_kind="worksheet_range",
        reference="A1:A1",
        has_header=False,
        columns=(TabularColumn(1, "coluna_1", None, "empty", True),),
        rows=(),
        formulas=(),
        warnings=warnings,
        merged_ranges=0,
        hidden_rows=0,
        hidden_columns=0,
    )


def _used_range(worksheet) -> str | None:
    nonempty = [
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if not isinstance(cell, MergedCell) and cell.value is not None
    ]
    if not nonempty:
        return None
    min_row = min(cell.row for cell in nonempty)
    max_row = max(cell.row for cell in nonempty)
    min_col = min(cell.column for cell in nonempty)
    max_col = max(cell.column for cell in nonempty)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _looks_like_header(values: list[object]) -> bool:
    if not values or any(value is None or not isinstance(value, str) for value in values):
        return False
    normalized = [value.strip().casefold() for value in values]
    return all(normalized) and len(set(normalized)) == len(normalized)


def _normalized_column_names(headers: list[str | None], width: int) -> list[str]:
    output: list[str] = []
    counts: dict[str, int] = {}
    for index in range(width):
        source = headers[index] or ""
        decomposed = unicodedata.normalize("NFKD", source)
        ascii_text = "".join(char for char in decomposed if not unicodedata.combining(char))
        base = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_text).strip("_").lower()
        base = base or f"coluna_{index + 1}"
        counts[base] = counts.get(base, 0) + 1
        output.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return output


def _normalize_cell_value(value: object, sanitizer: UnicodeSanitizer) -> TabularValue:
    if value is None or isinstance(value, (bool, int, date, datetime, time)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else sanitizer.clean(str(value))
    if isinstance(value, timedelta):
        return str(value)
    return sanitizer.clean(str(value))


def _infer_column_type(values: list[TabularValue]) -> str:
    types = {_value_type(value) for value in values if value is not None}
    if not types:
        return "empty"
    if types <= {"integer", "number"}:
        return "number" if "number" in types else "integer"
    if types <= {"date", "datetime"}:
        return "datetime" if "datetime" in types else "date"
    return next(iter(types)) if len(types) == 1 else "mixed"


def _value_type(value: TabularValue) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    return "string"


def _sheet_visibility(value: str) -> str:
    return "very_hidden" if value == "veryHidden" else value
